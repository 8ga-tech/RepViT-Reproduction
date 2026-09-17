"""自检 + 用 PowerPoint COM 把答辩 PPT 导出为 PDF（本机没有 LibreOffice，也没有 pywin32）。

用法::

    python tools/export_defense_pdf.py --check        # 只做 python-pptx 自检，不碰 PowerPoint
    python tools/export_defense_pdf.py                # 自检 + 导出 report/答辩PPT_RepViT.pdf

导出走 PowerShell 的 ``New-Object -ComObject PowerPoint.Application``（实测可用），
脚本退出前一定会 ``Quit()`` 并释放 COM 对象，导出前先清掉残留的 POWERPNT 进程。

自检项（对应验收口径）
----------------------
1. PPTX 能被 python-pptx 打开；总页数 ≤15；逐页打印文本/图片/图表/表格数量；
2. 五条曲线（train loss / val loss / val Top-1 / val Macro-F1 / lr）仍在；
3. 新增可视化齐全：每类 F1 柱状图、测试集 8 张预测、同图对比 4 组、训练集外实拍 5 张；
4. P03+P04 合并页存在，备注写明「合并」并逐条列出 15 个建议环节的落点；
5. 性能页（第 8 页训练曲线、第 14 页部署性能）元信息齐全：
   模型型号 / 输入尺寸 / batch size / 硬件 / 推理后端 / 精度类型 / 测试次数；
6. 旧写法（7.1e-06、完全等价、六个模型一致率 100%）已清除；
7. 导出后的 PDF 页数与 PPTX 一致。

自检报告写到 ``%TEMP%/repvit_ppt_selfcheck.md``（UTF-8）；控制台只打印 ASCII 摘要，
避免 PowerShell 控制台按 GBK 解码中文时报 UnicodeEncodeError。
"""
from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "report/答辩PPT_RepViT.pptx"
PDF = ROOT / "report/答辩PPT_RepViT.pdf"
REPORT = Path(tempfile.gettempdir()) / "repvit_ppt_selfcheck.md"
MAX_SLIDES = 15

# 性能图表元信息必须注明的字段（试题第 17 页）：任一候选命中即算该项齐全。
META_KEYS = {
    "模型型号": ["M0.9"],
    "输入尺寸": ["224×224", "1×3×224×224"],
    "batch size": ["batch 64", "batch 8", "batch 1"],
    "硬件": ["RTX 4060", "i7-13650HX"],
    "推理后端": ["PyTorch 2.14.0", "ONNX Runtime", "ORT "],
    "精度类型": ["FP32", "bf16"],
    "测试次数": ["评价次数", "预热 10", "每张只推理一次", "n=32"],
}
FORBIDDEN = ["7.1e-06", "完全等价", "6 个 ONNX 模型，PyTorch↔ONNX Top-1 一致率 100%"]
MILESTONES = ["1 题目与任务完成情况", "2 RepViT 核心结构", "3 为什么不是标准 ViT", "4 官方多型号评价",
              "5 数据集和训练流程", "6 Baseline 结果", "7 优化假设与控制变量", "8 曲线和定量结果",
              "9 混淆矩阵与失败案例", "10 Grad-CAM 结果", "11 结构重参数化", "12 ONNX 多模型部署",
              "13 性能比较", "14 遇到的问题", "15 总结"]


def iter_shapes(shape):
    yield shape
    if getattr(shape, "shapes", None) is not None:
        for child in shape.shapes:
            yield from iter_shapes(child)


def slide_text(slide):
    parts = []
    for sh in slide.shapes:
        for sub in iter_shapes(sh):
            if getattr(sub, "has_text_frame", False) and sub.text.strip():
                parts.append(sub.text)
            if getattr(sub, "has_table", False):
                for row in sub.table.rows:
                    parts.append(" / ".join(c.text for c in row.cells))
    return "\n".join(parts)


def slide_inventory(slide):
    inv = dict(text=0, picture=0, chart=0, table=0, charts=[], pics=[], texts=[])
    seen_tables = set()
    for sh in slide.shapes:
        for sub in iter_shapes(sh):
            if getattr(sub, "has_chart", False):
                inv["chart"] += 1
                inv["charts"].append([se.name for se in sub.chart.plots[0].series])
            elif getattr(sub, "has_table", False):
                if id(sub.table) not in seen_tables:
                    seen_tables.add(id(sub.table))
                    inv["table"] += 1
            elif sub.shape_type == 13:                      # MSO_SHAPE_TYPE.PICTURE
                inv["picture"] += 1
                inv["pics"].append(sub.image.sha1)          # 用 sha1 对齐源 PNG（filename 只是 image.png）
            elif getattr(sub, "has_text_frame", False):
                inv["text"] += 1
                if sub.text.strip():
                    inv["texts"].append(sub.text)
    return inv


def sha1_of(rel: str) -> str:
    import hashlib
    return hashlib.sha1((ROOT / rel).read_bytes()).hexdigest()


def check(prs, out):
    ok = True
    n = len(prs.slides)
    out.append(f"- 页数：{n}（上限 {MAX_SLIDES}）")
    if n > MAX_SLIDES:
        ok = False
        out.append("  - **失败**：页数超过 15 页")

    pages = []
    for i, s in enumerate(prs.slides, 1):
        inv = slide_inventory(s)
        pages.append(inv)
        out.append("")
        out.append(f"### P{i:02d}  文本 {inv['text']} / 图片 {inv['picture']} / 图表 {inv['chart']} / 表格 {inv['table']}")
        if inv["charts"]:
            out.append(f"- 图表系列：{inv['charts']}")
        if inv["pics"]:
            out.append(f"- 图片：{inv['pics']}")
        for t in inv["texts"]:
            out.append("  - " + t.replace("\n", " ⏎ ")[:240])

    whole = "\n".join(slide_text(s) for s in prs.slides)

    def need(label, cond, extra=""):
        nonlocal ok
        out.append(f"- {'通过' if cond else '**失败**'}：{label}{(' —— ' + extra) if extra else ''}")
        ok = ok and bool(cond)

    out += ["", "## 自检结论", ""]
    need("总页数 ≤15", n <= MAX_SLIDES, f"{n} 页")
    need("第 8 页仍是 5 条曲线（5 个折线图）", pages[7]["chart"] >= 5, f"图表数={pages[7]['chart']}")
    need("每类 F1 柱状图并入混淆矩阵页（第 9 页）",
         sha1_of("outputs/confusion_matrix/baseline_per_class_f1.png") in pages[8]["pics"],
         f"图片 {pages[8]['picture']} 张")
    need("混淆矩阵与 2 个失败案例仍在第 9 页",
         sha1_of("outputs/confusion_matrix/baseline_cm.png") in pages[8]["pics"]
         and sum(1 for p in pages[8]["pics"] if p in (sha1_of("outputs/predictions/case_wrong_baseline_case01.png"),
                                                      sha1_of("outputs/predictions/case_wrong_baseline_case02.png"))) == 2)
    need("新增「预测结果」页含 3 张预测拼图（grid8 / grid4 / grid5）",
         all(sha1_of(f"outputs/predictions/{n}.png") in pages[10]["pics"]
             for n in ("test_top5_baseline_grid8", "compare_baseline_vs_opt_combo_grid4",
                       "external_top5_pet37_grid5")),
         f"图片 {pages[10]['picture']} 张")
    need("P03+P04 合并为一页", "结构里没有注意力算子" in slide_text(prs.slides[2]))
    merged_notes = prs.slides[2].notes_slide.notes_text_frame.text if prs.slides[2].has_notes_slide else ""
    need("合并页备注写明合并事实", "合并页" in merged_notes and "环节 3" in merged_notes)
    missing_ms = [m for m in MILESTONES if m not in merged_notes]
    need("合并页备注列全 15 个建议环节", not missing_ms, f"缺 {missing_ms}" if missing_ms else "15/15")
    # 性能类页面都要 7 项元信息齐全（试题第 17 页）：P04 官方多型号评价表、
    # P08 训练曲线、P14 部署性能。P04 的口径行由 tools/update_defense.py 的
    # update_pretrained_meta() 补写（推理后端 / 测试次数）。
    for idx, label in ((3, "第 4 页官方多型号评价表"), (7, "第 8 页训练曲线"), (13, "第 14 页部署性能")):
        page = slide_text(prs.slides[idx])
        missing = [k for k, alts in META_KEYS.items() if not any(a in page for a in alts)]
        need(f"{label}元信息齐全（型号/输入尺寸/batch/硬件/后端/精度/测试次数）",
             not missing, f"缺 {missing}" if missing else "7/7")
    for bad in FORBIDDEN:
        need(f"旧写法已清除：{bad!r}", bad not in whole)
    need("重参数化与 ONNX 分开表述（7.093e-06 与 6.199e-06 各在自己的页）",
         "7.093e-06" in slide_text(prs.slides[11])
         and ("6.199e-06" in slide_text(prs.slides[12]) or "6.198883056640625e-06" in slide_text(prs.slides[12])))
    need("5.25e-06 只作为「未找到配套产物」的旧引用出现",
         whole.count("5.25e-06") >= 1 and "缺少对应完整记录" in whole)
    return ok


PS_TEMPLATE = r'''
$ErrorActionPreference = "Stop"
Get-Process POWERPNT -ErrorAction SilentlyContinue | Stop-Process -Force -ErrorAction SilentlyContinue
Start-Sleep -Milliseconds 800
$app = $null; $pres = $null
try {
    $app = New-Object -ComObject PowerPoint.Application
    $pres = $app.Presentations.Open("__PPTX__", $true, $false, $false)
    if (Test-Path "__PDF__") { Remove-Item "__PDF__" -Force }
    $pres.SaveAs("__PDF__", 32)
    Write-Output "SAVED"
} finally {
    if ($pres -ne $null) { $pres.Close() }
    if ($app -ne $null) { $app.Quit() }
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($pres) | Out-Null
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) | Out-Null
    [GC]::Collect()
}
'''


def export_pdf(pptx: Path, pdf: Path) -> bool:
    script = PS_TEMPLATE.replace("__PPTX__", str(pptx)).replace("__PDF__", str(pdf))
    tmp = Path(tempfile.gettempdir()) / "repvit_export_pdf.ps1"
    tmp.write_text(script, encoding="utf-8-sig")
    proc = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(tmp)],
                          capture_output=True, text=True)
    print(f"[pdf] powershell rc={proc.returncode}")
    if proc.stdout.strip():
        print("[pdf] stdout:", proc.stdout.strip()[-400:])
    if proc.stderr.strip():
        print("[pdf] stderr:", proc.stderr.strip()[-800:])
    return pdf.exists()


def pdf_pages(pdf: Path) -> int:
    import fitz
    with fitz.open(pdf) as doc:
        return doc.page_count


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="只自检，不导出 PDF")
    ap.add_argument("--pdf", default=str(PDF))
    a = ap.parse_args()

    prs = Presentation(PPTX)
    lines = ["# 答辩 PPT 自检报告", "", f"- 文件：`{PPTX.relative_to(ROOT).as_posix()}`", ""]
    ok = check(prs, lines)
    print(f"[check] slides={len(prs.slides)} checks_ok={ok} report={REPORT}")

    pdf = Path(a.pdf)
    if not a.check:
        saved = export_pdf(PPTX, pdf)
        if not saved:
            lines += ["", "## PDF", "- **失败**：PowerPoint COM 未生成 PDF"]
            ok = False
        else:
            n_pdf, n_ppt = pdf_pages(pdf), len(prs.slides)
            same = n_pdf == n_ppt
            lines += ["", "## PDF", f"- 页数：PDF {n_pdf} / PPTX {n_ppt} → {'一致' if same else '**不一致**'}"]
            print(f"[pdf] pages pdf={n_pdf} pptx={n_ppt} same={same}")
            ok = ok and same

    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
